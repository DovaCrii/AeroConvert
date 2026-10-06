"""`catalogos.py` con un ODBC de mentira: la lógica que rodea a la base, sin Access.

**Qué prueba esto y qué no.** El CI no tiene el controlador de Access (solo existe en Windows),
y sin él las funciones que abren la base —`esquema`, `a_excel`, `desde_excel`— quedaban sin una
sola prueba. Aquí `pyodbc` es un módulo falso que guarda las tablas en un archivo JSON. Se
comprueba lo que decide el código nuestro: qué se salta, cómo se nombran las hojas, qué se
rechaza, qué avisa, que un fallo a medias **deshaga todo** y que no quede un `.parcial`.

**No prueba que Access acepte lo que se escribe.** Para eso está el viaje completo marcado
`sin_access` en `test_catalogos.py`, que corre en la estación de trabajo. Un ODBC que yo mismo
escribí no es oráculo de nada de Access, y esta prueba no pretende serlo.
"""

from __future__ import annotations

import json
import re
import sys
import types
from pathlib import Path

import openpyxl
import pytest

from . import catalogos
from .composicion import ComposicionInvalida

CONTROLADOR = "Microsoft Access Driver (*.mdb, *.accdb)"


# --- El pyodbc de mentira --------------------------------------------------------------


class _Fila:
    def __init__(self, **campos):
        self.__dict__.update(campos)


class _Cursor:
    def __init__(self, conexion):
        self.c = conexion
        self._filas: list = []

    def tables(self, tableType=None):  # noqa: N803 - es el nombre de pyodbc
        return [_Fila(table_name=n) for n in self.c.datos["tablas"]]

    def columns(self, table):
        return [
            _Fila(column_name=n, type_name="VARCHAR", column_size=50)
            for n in self.c.datos["tablas"][table]["columnas"]
        ]

    def execute(self, sql, parametros=None):
        if (m := re.fullmatch(r"SELECT COUNT\(\*\) FROM \[(.+)\]", sql)) is not None:
            self._filas = [(len(self.c.datos["tablas"][m.group(1)]["filas"]),)]
        elif (m := re.fullmatch(r"SELECT \* FROM \[(.+)\]", sql)) is not None:
            self._filas = [tuple(f) for f in self.c.datos["tablas"][m.group(1)]["filas"]]
        elif (m := re.fullmatch(r"DELETE FROM \[(.+)\]", sql)) is not None:
            self.c.datos["tablas"][m.group(1)]["filas"] = []
        elif (m := re.fullmatch(r"INSERT INTO \[(.+?)\] \((.+)\) VALUES \(.+\)", sql)) is not None:
            if self.c.romper_al_insertar and parametros and self.c.romper_al_insertar in parametros:
                raise RuntimeError("('23000', '[23000] [Microsoft][ODBC] Valor duplicado (-1605)')")
            tabla = self.c.datos["tablas"][m.group(1)]
            campos = [c.strip(" []") for c in m.group(2).split(",")]
            tabla["filas"].append(
                [dict(zip(campos, parametros, strict=True)).get(n) for n in tabla["columnas"]]
            )
        else:  # pragma: no cover - una consulta que el código no debería emitir
            raise AssertionError(f"consulta no prevista: {sql}")
        return self

    def fetchone(self):
        return self._filas[0]

    def __iter__(self):
        return iter(self._filas)


class _Conexion:
    romper_al_insertar = None

    def __init__(self, ruta: Path):
        self.ruta = ruta
        self.datos = json.loads(ruta.read_text(encoding="utf-8"))
        self.cerrada = False

    def cursor(self):
        return _Cursor(self)

    def commit(self):
        self.ruta.write_text(json.dumps(self.datos), encoding="utf-8")

    def rollback(self):
        self.revertida = True

    def close(self):
        self.cerrada = True


def _modulo_falso(drivers=(CONTROLADOR,), romper_al_abrir=False, romper_al_insertar=None):
    modulo = types.ModuleType("pyodbc")
    modulo.drivers = lambda: list(drivers)
    conexiones: list[_Conexion] = []

    def connect(cadena, autocommit=True):
        if romper_al_abrir:
            raise RuntimeError("('HY000', \"[HY000] [Microsoft][ODBC] No es una base (-1028)\")")
        ruta = Path(re.search(r"DBQ=(.+?);", cadena).group(1))
        conexion = _Conexion(ruta)
        conexion.romper_al_insertar = romper_al_insertar
        conexiones.append(conexion)
        return conexion

    modulo.connect = connect
    modulo.conexiones = conexiones
    return modulo


def _base(ruta: Path, **tablas) -> Path:
    """Un «.mdb» de mentira: `nombre=(columnas, filas)`."""
    ruta.write_text(
        json.dumps(
            {
                "tablas": {
                    n: {"columnas": list(c), "filas": [list(f) for f in fs]}
                    for n, (c, fs) in tablas.items()
                }
            }
        ),
        encoding="utf-8",
    )
    return ruta


@pytest.fixture
def odbc(monkeypatch):
    modulo = _modulo_falso()
    monkeypatch.setitem(sys.modules, "pyodbc", modulo)
    monkeypatch.setenv("AEROCONVERT_CONTROLADOR_ACCESS", CONTROLADOR)
    return modulo


# --- La sonda y la cadena de conexión --------------------------------------------------


class TestLaSondaPorDentro:
    def test_sin_pyodbc_dice_que_falta_el_conector(self, monkeypatch):
        monkeypatch.setitem(sys.modules, "pyodbc", None)  # `import pyodbc` levanta ImportError
        estado = catalogos._mirar()
        assert not estado
        assert "conector" in estado.motivo
        assert estado.sugerencia

    def test_encuentra_el_controlador_de_mdb(self, monkeypatch):
        monkeypatch.setitem(
            sys.modules, "pyodbc", _modulo_falso(drivers=["SQL Server", CONTROLADOR])
        )
        assert catalogos._mirar().controlador == CONTROLADOR

    def test_sin_controlador_de_access_explica_como_instalarlo(self, monkeypatch):
        monkeypatch.setitem(sys.modules, "pyodbc", _modulo_falso(drivers=["SQL Server"]))
        estado = catalogos._mirar()
        assert not estado
        assert "Access Database Engine" in estado.sugerencia

    def test_la_sonda_se_acuerda_y_recordar_false_no(self, monkeypatch):
        monkeypatch.setitem(sys.modules, "pyodbc", _modulo_falso())
        from django.core.cache import cache

        cache.delete(catalogos.CLAVE_DE_CACHE)
        assert catalogos.sondar().controlador == CONTROLADOR
        monkeypatch.setitem(sys.modules, "pyodbc", _modulo_falso(drivers=[]))
        assert catalogos.sondar().controlador == CONTROLADOR, "la segunda vez sale de la caché"
        assert not catalogos.sondar(recordar=False)
        cache.delete(catalogos.CLAVE_DE_CACHE)

    def test_la_cadena_de_lectura_es_de_solo_lectura_y_la_de_crear_no(self, odbc, tmp_path):
        assert "ReadOnly=1" in catalogos._cadena(tmp_path / "a.mdb")
        assert "ReadOnly" not in catalogos._cadena(tmp_path / "a.mdb", crear=True)

    def test_sin_controlador_la_cadena_se_niega_con_el_motivo(self, monkeypatch, tmp_path):
        monkeypatch.delenv("AEROCONVERT_CONTROLADOR_ACCESS", raising=False)
        monkeypatch.setitem(sys.modules, "pyodbc", _modulo_falso(drivers=[]))
        from django.core.cache import cache

        cache.delete(catalogos.CLAVE_DE_CACHE)
        with pytest.raises(ComposicionInvalida, match="controlador"):
            catalogos._cadena(tmp_path / "a.mdb")
        cache.delete(catalogos.CLAVE_DE_CACHE)


# --- El esquema ----------------------------------------------------------------------


class TestEsquema:
    def test_cuenta_filas_y_se_salta_las_tablas_del_sistema(self, odbc, tmp_path):
        base = _base(
            tmp_path / "c.mdb",
            PIPE=(["MATERIAL", "OD"], [["A", 1], ["B", 2]]),
            MSysObjects=(["Id"], [[1]]),
        )
        tablas = catalogos.esquema(base)
        assert [(t.nombre, t.filas) for t in tablas] == [("PIPE", 2)]
        assert [c.nombre for c in tablas[0].columnas] == ["MATERIAL", "OD"]
        assert odbc.conexiones[-1].cerrada

    def test_un_archivo_que_no_esta(self, odbc, tmp_path):
        with pytest.raises(ComposicionInvalida, match="No está"):
            catalogos.esquema(tmp_path / "no-existe.mdb")

    def test_una_base_que_no_abre_dice_por_que_sin_el_ruido_de_odbc(self, monkeypatch, tmp_path):
        monkeypatch.setitem(sys.modules, "pyodbc", _modulo_falso(romper_al_abrir=True))
        monkeypatch.setenv("AEROCONVERT_CONTROLADOR_ACCESS", CONTROLADOR)
        base = _base(tmp_path / "c.mdb", PIPE=(["A"], []))
        with pytest.raises(ComposicionInvalida) as fallo:
            catalogos.esquema(base)
        assert "No se pudo abrir c.mdb" in str(fallo.value)
        assert "HY000" not in str(fallo.value)

    def test_sin_tablas_de_usuario_no_hay_nada_que_exportar(self, odbc, tmp_path):
        base = _base(tmp_path / "c.mdb", MSysObjects=(["Id"], []))
        with pytest.raises(ComposicionInvalida, match="ninguna tabla"):
            catalogos.esquema(base)

    def test_un_nombre_de_tabla_que_se_sale_del_corchete_se_para(self, odbc, tmp_path):
        base = _base(tmp_path / "c.mdb", **{"PIPE]; DROP TABLE x;--": (["A"], [])})
        with pytest.raises(ComposicionInvalida, match="nombre que no se puede usar"):
            catalogos.esquema(base)


# --- A Excel -------------------------------------------------------------------------


class TestAExcel:
    def test_una_hoja_por_tabla_con_encabezados_y_la_fila_fija(self, odbc, tmp_path):
        base = _base(
            tmp_path / "c.mdb",
            PIPE=(["MATERIAL", "OD"], [["acero", 10], ["pvc", 20]]),
            ELBOW=(["TIPO"], []),
        )
        salida = catalogos.a_excel(base, destino=tmp_path / "c.xlsx")

        libro = openpyxl.load_workbook(salida)
        assert libro.sheetnames == ["PIPE", "ELBOW"]
        assert [[c.value for c in f] for f in libro["PIPE"].iter_rows()] == [
            ["MATERIAL", "OD"],
            ["acero", 10],
            ["pvc", 20],
        ]
        assert libro["ELBOW"].max_row == 1, "una tabla vacía sale igual, con sus encabezados"
        assert libro["PIPE"].freeze_panes == "A2"
        assert not list(tmp_path.glob("*.parcial")), "no queda el parcial"

    def test_sin_destino_lo_deja_junto_al_origen(self, odbc, tmp_path):
        base = _base(tmp_path / "c.mdb", PIPE=(["A"], [[1]]))
        assert catalogos.a_excel(base) == tmp_path / "c.xlsx"

    def test_un_campo_binario_se_dice_cuantos_bytes_es(self, odbc, tmp_path):
        assert catalogos._a_celda(b"abc") == "<3 bytes>"
        assert catalogos._a_celda(bytearray(5)) == "<5 bytes>"
        assert catalogos._a_celda("x") == "x"

    def test_pasarse_del_tope_corta_y_no_deja_archivo(self, odbc, tmp_path, monkeypatch):
        monkeypatch.setattr(catalogos, "TOPE_FILAS", 2)
        base = _base(tmp_path / "c.mdb", PIPE=(["A"], [[1], [2], [3]]))
        with pytest.raises(ComposicionInvalida, match="pasa de 2 filas"):
            catalogos.a_excel(base, destino=tmp_path / "c.xlsx")
        assert not (tmp_path / "c.xlsx").exists()
        assert odbc.conexiones[-1].cerrada

    def test_los_nombres_largos_o_repetidos_no_chocan_en_excel(self):
        usados: set[str] = set()
        largo = "T" * 40
        primero = catalogos._nombre_de_hoja(largo, usados)
        segundo = catalogos._nombre_de_hoja(largo, usados)
        assert len(primero) == 31 and len(segundo) == 31 and primero != segundo
        assert segundo.endswith("_2")
        assert catalogos._nombre_de_hoja("a:b*c?", usados) == "a_b_c_"
        assert catalogos._nombre_de_hoja(":::", set()) == "___"
        assert catalogos._nombre_de_hoja("", set()) == "Tabla"
        assert catalogos._nombre_de_hoja("PIPE", {"pipe"}) == "PIPE_2", (
            "Excel no distingue mayúsculas"
        )


# --- De vuelta a Access ----------------------------------------------------------------


def _excel(ruta: Path, **hojas) -> Path:
    libro = openpyxl.Workbook()
    libro.remove(libro.active)
    for nombre, filas in hojas.items():
        hoja = libro.create_sheet(nombre)
        for fila in filas:
            hoja.append(list(fila))
    libro.save(ruta)
    return ruta


def _leer(ruta: Path) -> dict:
    return json.loads(ruta.read_text(encoding="utf-8"))["tablas"]


class TestDesdeExcel:
    def test_sustituye_las_filas_de_la_plantilla_y_la_deja_intacta(self, odbc, tmp_path):
        plantilla = _base(tmp_path / "plantilla.mdb", PIPE=(["MATERIAL", "OD"], [["viejo", 0]]))
        antes = plantilla.read_bytes()
        hoja = _excel(tmp_path / "h.xlsx", PIPE=[("MATERIAL", "OD"), ("acero", 10), ("pvc", 20)])

        destino, avisos = catalogos.desde_excel(hoja, plantilla, destino=tmp_path / "nuevo.mdb")

        assert _leer(destino)["PIPE"]["filas"] == [["acero", 10], ["pvc", 20]]
        assert plantilla.read_bytes() == antes, "la plantilla no se toca"
        assert avisos == []
        assert not list(tmp_path.glob("*.parcial"))

    def test_las_filas_en_blanco_del_final_de_la_hoja_no_son_datos(self, odbc, tmp_path):
        plantilla = _base(tmp_path / "p.mdb", PIPE=(["MATERIAL"], []))
        hoja = _excel(tmp_path / "h.xlsx", PIPE=[("MATERIAL",), ("acero",), (None,), ("  ",)])
        destino, _ = catalogos.desde_excel(hoja, plantilla, destino=tmp_path / "n.mdb")
        assert _leer(destino)["PIPE"]["filas"] == [["acero"]]

    def test_una_tabla_que_el_excel_no_trae_queda_vacia_y_se_avisa(self, odbc, tmp_path):
        plantilla = _base(tmp_path / "p.mdb", PIPE=(["A"], [["x"]]), ELBOW=(["B"], [["y"], ["z"]]))
        hoja = _excel(tmp_path / "h.xlsx", PIPE=[("A",), ("n",)])
        destino, avisos = catalogos.desde_excel(hoja, plantilla, destino=tmp_path / "n.mdb")
        assert _leer(destino)["ELBOW"]["filas"] == []
        assert any("ELBOW no venía en el Excel" in a for a in avisos)

    def test_una_hoja_sin_filas_de_datos_avisa(self, odbc, tmp_path):
        plantilla = _base(tmp_path / "p.mdb", PIPE=(["A"], [["x"]]))
        hoja = _excel(tmp_path / "h.xlsx", PIPE=[("A",)])
        destino, avisos = catalogos.desde_excel(hoja, plantilla, destino=tmp_path / "n.mdb")
        assert _leer(destino)["PIPE"]["filas"] == []
        assert any("no traía ninguna fila" in a for a in avisos)

    def test_una_hoja_completamente_vacia_se_salta(self, odbc, tmp_path):
        plantilla = _base(tmp_path / "p.mdb", PIPE=(["A"], [["x"]]))
        libro = openpyxl.Workbook()
        libro.active.title = "PIPE"
        libro.save(tmp_path / "h.xlsx")
        destino, avisos = catalogos.desde_excel(
            tmp_path / "h.xlsx", plantilla, destino=tmp_path / "n.mdb"
        )
        assert _leer(destino)["PIPE"]["filas"] == [], "la tabla queda vacía y se dice"
        assert avisos

    def test_sin_plantilla_no_hay_nada_que_copiar(self, odbc, tmp_path):
        hoja = _excel(tmp_path / "h.xlsx", PIPE=[("A",)])
        with pytest.raises(ComposicionInvalida, match="No está la plantilla"):
            catalogos.desde_excel(hoja, tmp_path / "no.mdb")

    def test_una_hoja_desconocida_se_rechaza_antes_de_escribir(self, odbc, tmp_path):
        plantilla = _base(tmp_path / "p.mdb", PIPE=(["A"], []))
        hoja = _excel(tmp_path / "h.xlsx", Otra=[("A",)])
        with pytest.raises(ComposicionInvalida, match="hojas que la plantilla no tiene"):
            catalogos.desde_excel(hoja, plantilla, destino=tmp_path / "n.mdb")
        assert not (tmp_path / "n.mdb").exists()
        assert not list(tmp_path.glob("*.parcial"))

    def test_un_fallo_a_mitad_deshace_todo_y_no_deja_parcial(self, tmp_path, monkeypatch):
        modulo = _modulo_falso(romper_al_insertar="duplicado")
        monkeypatch.setitem(sys.modules, "pyodbc", modulo)
        monkeypatch.setenv("AEROCONVERT_CONTROLADOR_ACCESS", CONTROLADOR)
        plantilla = _base(tmp_path / "p.mdb", PIPE=(["A"], [["viejo"]]))
        hoja = _excel(tmp_path / "h.xlsx", PIPE=[("A",), ("bueno",), ("duplicado",)])

        with pytest.raises(ComposicionInvalida) as fallo:
            catalogos.desde_excel(hoja, plantilla, destino=tmp_path / "n.mdb")

        assert "PIPE, fila 3 del Excel" in str(fallo.value)
        assert "Valor duplicado" in str(fallo.value) and "23000" not in str(fallo.value)
        assert not (tmp_path / "n.mdb").exists()
        assert not list(tmp_path.glob("*.parcial"))
        assert modulo.conexiones[-1].revertida and modulo.conexiones[-1].cerrada
        assert _leer(plantilla)["PIPE"]["filas"] == [["viejo"]]

    def test_pasarse_del_tope_de_filas_corta(self, odbc, tmp_path, monkeypatch):
        monkeypatch.setattr(catalogos, "TOPE_FILAS", 1)
        plantilla = _base(tmp_path / "p.mdb", PIPE=(["A"], []))
        hoja = _excel(tmp_path / "h.xlsx", PIPE=[("A",), ("1",), ("2",)])
        with pytest.raises(ComposicionInvalida, match="pasa de 1 filas"):
            catalogos.desde_excel(hoja, plantilla, destino=tmp_path / "n.mdb")
        assert not (tmp_path / "n.mdb").exists()

    def test_una_hoja_sin_cabeceras_no_inserta_nada(self, odbc, tmp_path):
        plantilla = _base(tmp_path / "p.mdb", PIPE=(["A"], []))
        hoja = _excel(tmp_path / "h.xlsx", PIPE=[(None, None), ("x", "y")])
        destino, avisos = catalogos.desde_excel(hoja, plantilla, destino=tmp_path / "n.mdb")
        assert _leer(destino)["PIPE"]["filas"] == []


class TestLimpiarElMensajeDeOdbc:
    def test_se_queda_con_la_frase_util(self):
        feo = RuntimeError(
            "('HY000', \"[HY000] [Microsoft][Controlador ODBC] No se puede abrir (-1811)\")"
        )
        assert catalogos._limpiar(feo) == "No se puede abrir"

    def test_un_mensaje_sin_esa_forma_sale_tal_cual(self):
        assert catalogos._limpiar(RuntimeError("sin forma")) == "sin forma"
