"""«Office a PDF» con LibreOffice cuando no hay Office (F17.1).

El hijo se prueba con un **`soffice` de mentira** que hace lo que hace el de verdad: escribe
`<outdir>/<nombre>.pdf` y deja un archivo de bloqueo junto al documento que abre. Así se comprueba
el pegamento —que el original y su carpeta no se tocan, que el PDF llega al parcial, que la carpeta
de trabajo se borra siempre— sin LibreOffice en la máquina. La conversión de verdad se mide donde
esté instalado (`@pytest.mark.oraculo`, abajo).
"""

from __future__ import annotations

import hashlib
import io
import os
import shutil
import subprocess
import sys
from pathlib import Path

import pypdfium2
import pytest
from django.contrib.auth import get_user_model
from django.urls import reverse

from apps.documents import libreoffice_hijo, motor, office
from apps.documents.views import _comun

FALSO = r"""
import sys, time
from pathlib import Path
from reportlab.pdfgen import canvas
args = sys.argv[1:]
modo = Path(__file__).with_suffix(".modo").read_text().strip()
outdir = Path(args[args.index("--outdir") + 1])
entrada = Path(args[-1])
(entrada.parent / (".~lock." + entrada.name + "#")).write_text("bloqueo")
if modo == "duerme":
    time.sleep(30)
if modo == "bien":
    c = canvas.Canvas(str(outdir / (entrada.stem + ".pdf")))
    c.drawString(100, 700, "Texto del documento de partida")
    c.showPage(); c.save()
if modo == "vacio":
    (outdir / (entrada.stem + ".pdf")).write_bytes(b"")
"""


def _soffice_falso(carpeta: Path, modo: str) -> str:
    guion = carpeta / "soffice_falso.py"
    guion.write_text(FALSO, encoding="utf-8")
    guion.with_suffix(".modo").write_text(modo, encoding="utf-8")
    if os.name == "nt":
        lanzador = carpeta / "soffice_falso.cmd"
        lanzador.write_text(f'@"{sys.executable}" "{guion}" %*\n', encoding="utf-8")
    else:
        lanzador = carpeta / "soffice_falso"
        lanzador.write_text(
            f'#!/bin/sh\nexec "{sys.executable}" "{guion}" "$@"\n', encoding="utf-8"
        )
        lanzador.chmod(0o755)
    return str(lanzador)


def _documento(carpeta: Path) -> Path:
    compartida = carpeta / "compartida"
    compartida.mkdir()
    doc = compartida / "Informe final.docx"
    doc.write_bytes(b"PK\x03\x04 no hace falta que sea un docx de verdad")
    return doc


def _huellas(carpeta: Path):
    return {
        p.name: (hashlib.sha256(p.read_bytes()).hexdigest(), p.stat().st_mtime_ns)
        for p in carpeta.iterdir()
    }


class TestElHijo:
    def test_convierte_y_deja_el_pdf_en_el_parcial_sin_tocar_la_carpeta_del_original(
        self, tmp_path
    ):
        doc = _documento(tmp_path)
        antes = _huellas(doc.parent)
        parcial = tmp_path / "salida.parcial.pdf"
        libreoffice_hijo.convertir(_soffice_falso(tmp_path, "bien"), doc, parcial, 60)
        assert len(pypdfium2.PdfDocument(str(parcial))) == 1
        texto = pypdfium2.PdfDocument(str(parcial))[0].get_textpage().get_text_bounded()
        assert "Texto del documento de partida" in texto
        assert _huellas(doc.parent) == antes, "ni el original ni su carpeta (ni un .~lock)"
        assert not parcial.with_name(parcial.name + ".lo").exists(), "la carpeta de trabajo se va"

    @pytest.mark.parametrize("modo", ["nada", "vacio"])
    def test_si_no_escribe_el_pdf_falla_aunque_salga_con_cero(self, tmp_path, modo):
        doc = _documento(tmp_path)
        antes = _huellas(doc.parent)
        parcial = tmp_path / "salida.parcial.pdf"
        with pytest.raises(SystemExit) as salida:
            libreoffice_hijo.convertir(_soffice_falso(tmp_path, modo), doc, parcial, 60)
        assert salida.value.code == 2
        assert not parcial.exists() and not parcial.with_name(parcial.name + ".lo").exists()
        assert _huellas(doc.parent) == antes

    def test_si_se_cuelga_se_corta_y_se_limpia(self, tmp_path):
        doc = _documento(tmp_path)
        parcial = tmp_path / "salida.parcial.pdf"
        with pytest.raises(SystemExit) as salida:
            libreoffice_hijo.convertir(_soffice_falso(tmp_path, "duerme"), doc, parcial, 2)
        assert salida.value.code == 3
        assert not parcial.with_name(parcial.name + ".lo").exists()

    def test_cada_trabajo_lleva_su_propio_perfil(self, tmp_path, monkeypatch):
        vistos = []

        def falso_run(orden, **kw):
            vistos.append(orden)
            salida = Path(orden[orden.index("--outdir") + 1])
            (salida / "documento.pdf").write_bytes(b"%PDF-1.4 x")
            return subprocess.CompletedProcess(orden, 0, "", "")

        monkeypatch.setattr(libreoffice_hijo.subprocess, "run", falso_run)
        doc = _documento(tmp_path)
        libreoffice_hijo.convertir("soffice", doc, tmp_path / "a.parcial.pdf", 60)
        libreoffice_hijo.convertir("soffice", doc, tmp_path / "b.parcial.pdf", 60)
        perfiles = [next(a for a in o if a.startswith("-env:UserInstallation=")) for o in vistos]
        assert perfiles[0] != perfiles[1]
        assert all("--headless" in o and "--convert-to" in o for o in vistos)


@pytest.fixture
def sin_office_con_libre(monkeypatch):
    monkeypatch.setattr(
        office, "sondar", lambda **kw: office.Disponible(frozenset(), motivo="No hay Office.")
    )
    monkeypatch.setattr(office, "sondar_libreoffice", lambda **kw: "/usr/bin/soffice")


@pytest.fixture
def sin_nada(monkeypatch):
    monkeypatch.setattr(
        office, "sondar", lambda **kw: office.Disponible(frozenset(), motivo="No hay Office.")
    )
    monkeypatch.setattr(office, "sondar_libreoffice", lambda **kw: "")


def test_si_el_hijo_muere_el_corredor_borra_su_carpeta_de_trabajo(tmp_path):
    from apps.jobs.runner import _borrar

    parcial = tmp_path / "x.parcial.pdf"
    trabajo = parcial.with_name(parcial.name + ".lo")
    (trabajo / "perfil").mkdir(parents=True)
    (trabajo / "perfil" / "algo").write_text("x")
    ajena = tmp_path / "x.parcial.pdf.respaldo"
    (ajena / "dentro").mkdir(parents=True)
    _borrar(parcial)
    assert not trabajo.exists()
    assert (ajena / "dentro").exists(), "una carpeta que no es la de LibreOffice no se toca"


class TestLaDisponibilidad:
    def test_con_libreoffice_solo_si_el_trabajo_lo_pidio(self, sin_office_con_libre):
        pedido = motor.disponibilidad("office", {"con_libreoffice": True})
        assert pedido.disponible and pedido.version == motor.ROTULO_LIBREOFFICE
        # Un lote o la API, que no pueden aceptar la advertencia: apagada, con su código.
        sin_pedir = motor.disponibilidad("office")
        assert not sin_pedir.disponible and sin_pedir.codigo_motivo == "sin-office"
        assert not motor.disponibilidad("a_word").disponible, "PDF a Word sigue siendo de Word"

    def test_el_recibo_dice_libreoffice_aunque_haya_office(self, monkeypatch):
        """Encolado con LibreOffice y con Office instalado después: se convierte con lo pedido y
        el recibo lo dice."""
        monkeypatch.setattr(office, "sondar", lambda **kw: office.Disponible(frozenset({"word"})))
        monkeypatch.setattr(office, "sondar_libreoffice", lambda **kw: "/usr/bin/soffice")
        assert motor.disponibilidad("office", {"con_libreoffice": True}).version == (
            motor.ROTULO_LIBREOFFICE
        )
        assert motor.disponibilidad("office").version == "documentos:office"

    def test_si_se_desinstalo_entre_encolar_y_ejecutar_falla_con_su_codigo(self, sin_nada):
        estado = motor.disponibilidad("office", {"con_libreoffice": True})
        assert not estado.disponible and estado.codigo_motivo == "sin-office"

    def test_sin_ninguno_sale_apagado_con_su_motivo(self, sin_nada):
        estado = motor.disponibilidad("office")
        assert not estado.disponible and estado.codigo_motivo == "sin-office"

    def test_el_indice_dice_con_que_y_que_puede_variar(self, sin_office_con_libre):
        filas = {f["id"]: f for f in _comun.estado_de_herramientas()}
        assert filas["office"]["disponible"] and "puede variar" in filas["office"]["variante"]
        assert not filas["a_word"]["disponible"]


@pytest.fixture
def entrado(client, db, settings, tmp_path):
    settings.RAICES_PERMITIDAS = str(tmp_path)
    settings.CARPETA_DE_TRABAJO = str(tmp_path / "trabajo")
    client.force_login(get_user_model().objects.create_user("ana", password="x" * 20))  # nosec B106
    return client


class TestLaPantalla:
    def test_sin_office_ofrece_libreoffice_con_la_advertencia(self, entrado, sin_office_con_libre):
        cuerpo = entrado.get(reverse("documents:office")).content.decode()
        assert 'name="con_libreoffice"' in cuerpo and "puede variar" in cuerpo

    def test_sin_aceptar_no_se_encola_nada(self, entrado, sin_office_con_libre, tmp_path):
        from apps.jobs.models import ConversionJob

        doc = _documento(tmp_path)
        respuesta = entrado.post(reverse("documents:office"), {"ruta": str(doc)})
        assert respuesta.status_code == 200 and "Acepto" in respuesta.content.decode()
        assert not ConversionJob.objects.exists()

    def test_aceptando_se_encola_con_libreoffice_y_el_plan_lo_usa(
        self, entrado, sin_office_con_libre, tmp_path
    ):
        from apps.jobs.models import ConversionJob

        doc = _documento(tmp_path)
        respuesta = entrado.post(
            reverse("documents:office"), {"ruta": str(doc), "con_libreoffice": "si"}
        )
        assert respuesta.status_code == 302
        trabajo = ConversionJob.objects.get()
        assert trabajo.options == {"con_libreoffice": True}
        argv = motor.plan(trabajo).argv
        assert "apps.documents.libreoffice_hijo" in argv and "/usr/bin/soffice" in argv

    def test_de_extremo_a_extremo_el_corredor_rotula_y_el_original_queda(
        self, entrado, monkeypatch, tmp_path
    ):
        """Por la cola, con el `soffice` de mentira: hecho, rotulado y sin tocar el original."""
        from apps.jobs import despachador
        from apps.jobs.models import ConversionJob

        monkeypatch.setattr(
            office, "sondar", lambda **kw: office.Disponible(frozenset(), motivo="No hay Office.")
        )
        falso = _soffice_falso(tmp_path, "bien")
        monkeypatch.setattr(office, "sondar_libreoffice", lambda **kw: falso)
        doc = _documento(tmp_path)
        antes = _huellas(doc.parent)
        entrado.post(reverse("documents:office"), {"ruta": str(doc), "con_libreoffice": "si"})
        despachador.procesar_una_vez()
        trabajo = ConversionJob.objects.get()
        assert trabajo.status == "done", trabajo.reason_detail
        assert trabajo.engine_version == motor.ROTULO_LIBREOFFICE
        assert len(pypdfium2.PdfDocument(trabajo.output_path)) == 1
        despues = _huellas(doc.parent)
        assert {k: v for k, v in despues.items() if k in antes} == antes

    def test_un_fallo_por_la_cola_deja_el_original_y_no_deja_restos(
        self, entrado, monkeypatch, tmp_path
    ):
        from apps.jobs import despachador
        from apps.jobs.models import ConversionJob

        monkeypatch.setattr(
            office, "sondar", lambda **kw: office.Disponible(frozenset(), motivo="No hay Office.")
        )
        falso = _soffice_falso(tmp_path, "nada")
        monkeypatch.setattr(office, "sondar_libreoffice", lambda **kw: falso)
        doc = _documento(tmp_path)
        antes = _huellas(doc.parent)
        entrado.post(reverse("documents:office"), {"ruta": str(doc), "con_libreoffice": "si"})
        despachador.procesar_una_vez()
        trabajo = ConversionJob.objects.get()
        assert trabajo.status == "error"
        assert _huellas(doc.parent) == antes
        assert not list(Path(trabajo.output_path).parent.glob("*.parcial*"))


@pytest.mark.oraculo
@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice no está en esta máquina")
def test_libreoffice_de_verdad_convierte_un_docx(tmp_path):
    """Con LibreOffice instalado: PDFium cuenta las páginas y halla el texto del documento."""
    import docx

    documento = docx.Document()
    documento.add_paragraph("Frase de control 4711 para el oráculo")
    memoria = io.BytesIO()
    documento.save(memoria)
    carpeta = tmp_path / "c"
    carpeta.mkdir()
    origen = carpeta / "control.docx"
    origen.write_bytes(memoria.getvalue())
    parcial = tmp_path / "s.parcial.pdf"
    libreoffice_hijo.convertir(shutil.which("soffice"), origen, parcial, 120)
    pdf = pypdfium2.PdfDocument(str(parcial))
    assert len(pdf) >= 1
    assert "Frase de control 4711" in pdf[0].get_textpage().get_text_bounded()
