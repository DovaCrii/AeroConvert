"""Leer un catálogo de Access con `mdbtools`, sin el motor de Microsoft (F17.2).

En la estación hay ACE y no hay mdbtools; en el CI, ninguno de los dos. Así que se prueba con un
`mdbtools` **de mentira** que contesta como el de verdad (`mdb-tables -1`, `mdb-export` con su CSV,
`mdb-json` una fila por línea) sobre un catálogo descrito aquí. El Excel resultante se lee con
openpyxl y se compara con lo que se describió. La lectura con el `mdbtools` real, contra las cifras
que dio ACE en la estación, se hace en `p340` (ver `docs/PRUEBAS_CON_ORACULO.md`).
"""

from __future__ import annotations

import json
import os
import sys
from pathlib import Path

import openpyxl
import pytest
from django.contrib.auth import get_user_model
from django.urls import reverse

from apps.documents import catalogos, motor
from apps.documents.composicion import ComposicionInvalida
from apps.documents.tarea import VARIABLE_ACCESS

CATALOGO = {
    "PIPE": {
        "columnas": ["PartFamilyId", "NOMINAL_DIAMETER", "WALL_THICKNESS", "PIECE_MARK"],
        "filas": [[1, 110.0, 6.6, "P-110"], [2, 160.0, 9.5, "P-160"], [3, 200.0, 11.9, "P-200"]],
    },
    "ELBOW": {"columnas": ["PartFamilyId", "ANGLE"], "filas": [[10, 90], [11, 45]]},
    "MSysObjects": {"columnas": ["Id"], "filas": [[1]]},
}

FALSO = r"""
import csv, json, sys
from pathlib import Path
catalogo = json.loads(Path(__file__).with_name("catalogo.json").read_text())
programa = Path(__file__).stem
args = sys.argv[1:]
if programa == "mdb-tables":
    print("\n".join(catalogo))
elif programa == "mdb-export":
    tabla = catalogo[args[1]]
    escritor = csv.writer(sys.stdout, lineterminator="\n")
    escritor.writerow(tabla["columnas"])
    for fila in tabla["filas"]:
        escritor.writerow(fila)
elif programa == "mdb-json":
    tabla = catalogo[args[1]]
    for fila in tabla["filas"]:
        print(json.dumps(dict(zip(tabla["columnas"], fila))))
"""


def _mdbtools(carpeta: Path, catalogo=CATALOGO, *, con_json=True) -> str:
    carpeta.mkdir(parents=True, exist_ok=True)
    (carpeta / "catalogo.json").write_text(json.dumps(catalogo), encoding="utf-8")
    programas = ["mdb-tables", "mdb-export"] + (["mdb-json"] if con_json else [])
    for programa in programas:
        guion = carpeta / f"{programa}.py"
        guion.write_text(FALSO, encoding="utf-8")
        if os.name == "nt":
            (carpeta / f"{programa}.cmd").write_text(
                f'@"{sys.executable}" "{guion}" %*\n', encoding="utf-8"
            )
        else:
            lanzador = carpeta / programa
            lanzador.write_text(f'#!/bin/sh\nexec "{sys.executable}" "{guion}" "$@"\n')
            lanzador.chmod(0o755)
    return str(carpeta)


@pytest.fixture
def con_mdbtools(tmp_path, monkeypatch):
    carpeta = _mdbtools(tmp_path / "mdbtools")
    monkeypatch.setenv(VARIABLE_ACCESS, f"{catalogos.PREFIJO_MDBTOOLS}{carpeta}")
    base = tmp_path / "HDPE.mdb"
    base.write_bytes(b"\x00\x01\x00\x00Standard Jet DB\x00 una base de mentira")
    return base


class TestLaLectura:
    def test_el_esquema_salta_las_tablas_del_sistema_y_cuenta_las_filas(self, con_mdbtools):
        tablas = {t.nombre: t for t in catalogos.esquema(con_mdbtools)}
        assert set(tablas) == {"PIPE", "ELBOW"}
        assert tablas["PIPE"].filas == 3 and tablas["ELBOW"].filas == 2
        assert [c.nombre for c in tablas["PIPE"].columnas] == CATALOGO["PIPE"]["columnas"]
        assert all(c.tipo == "?" for c in tablas["PIPE"].columnas), "el tipo no se inventa"

    def test_el_excel_trae_una_hoja_por_tabla_con_los_valores_y_sus_tipos(
        self, con_mdbtools, tmp_path
    ):
        destino = catalogos.a_excel(con_mdbtools, destino=tmp_path / "catalogo.xlsx")
        libro = openpyxl.load_workbook(destino, read_only=True)
        assert libro.sheetnames == ["PIPE", "ELBOW"]
        filas = list(libro["PIPE"].iter_rows(values_only=True))
        assert list(filas[0]) == CATALOGO["PIPE"]["columnas"]
        assert [list(f) for f in filas[1:]] == CATALOGO["PIPE"]["filas"]
        assert isinstance(filas[1][2], float), "con mdb-json, los números van como números"

    def test_sin_mdb_json_van_como_texto(self, tmp_path, monkeypatch):
        carpeta = _mdbtools(tmp_path / "viejo", con_json=False)
        monkeypatch.setenv(VARIABLE_ACCESS, f"{catalogos.PREFIJO_MDBTOOLS}{carpeta}")
        monkeypatch.setattr("shutil.which", lambda nombre: None)
        base = tmp_path / "b.mdb"
        base.write_bytes(b"x")
        destino = catalogos.a_excel(base, destino=tmp_path / "c.xlsx")
        filas = list(
            openpyxl.load_workbook(destino, read_only=True)["ELBOW"].iter_rows(values_only=True)
        )
        assert filas[1] == ("10", "90")

    def test_un_nombre_de_tabla_hostil_se_rechaza(self, tmp_path, monkeypatch):
        carpeta = _mdbtools(tmp_path / "h", {"PIPE]; DROP": {"columnas": ["a"], "filas": []}})
        monkeypatch.setenv(VARIABLE_ACCESS, f"{catalogos.PREFIJO_MDBTOOLS}{carpeta}")
        base = tmp_path / "b.mdb"
        base.write_bytes(b"x")
        with pytest.raises(ComposicionInvalida, match="nombre"):
            catalogos.esquema(base)

    def test_la_base_no_se_toca(self, con_mdbtools, tmp_path):
        import hashlib

        huella = (
            hashlib.sha256(con_mdbtools.read_bytes()).hexdigest(),
            con_mdbtools.stat().st_mtime_ns,
        )
        catalogos.a_excel(con_mdbtools, destino=tmp_path / "c.xlsx")
        assert (
            hashlib.sha256(con_mdbtools.read_bytes()).hexdigest(),
            con_mdbtools.stat().st_mtime_ns,
        ) == huella


class TestLaSonda:
    def test_con_ace_manda_ace(self, monkeypatch):
        monkeypatch.setattr(
            catalogos, "sondar", lambda **k: catalogos.Disponible(controlador="ACE")
        )
        assert catalogos.sondar_lectura().controlador == "ACE"

    def test_sin_ace_y_con_mdbtools_se_puede_leer_pero_no_escribir(
        self, monkeypatch, tmp_path, settings
    ):
        monkeypatch.setattr(catalogos, "sondar", lambda **k: catalogos.Disponible(motivo="sin ACE"))
        settings.MDBTOOLS = _mdbtools(tmp_path / "m")
        (Path(settings.MDBTOOLS) / "mdb-tables").write_text("x")  # lo que busca `encontrar`
        assert catalogos.sondar_lectura().controlador.startswith(catalogos.PREFIJO_MDBTOOLS)
        assert motor.disponibilidad("catalogo_excel").disponible
        assert not motor.disponibilidad("excel_catalogo").disponible, "escribir sigue pidiendo ACE"

    def test_sin_ninguno_sale_apagada_con_sin_access(self, monkeypatch, settings):
        monkeypatch.setattr(catalogos, "sondar", lambda **k: catalogos.Disponible(motivo="sin ACE"))
        monkeypatch.setattr("shutil.which", lambda nombre: None)
        settings.MDBTOOLS = ""
        estado = motor.disponibilidad("catalogo_excel")
        assert not estado.disponible and estado.codigo_motivo == "sin-access"


def test_la_pantalla_lo_ofrece_con_mdbtools(client, db, monkeypatch, tmp_path):
    client.force_login(get_user_model().objects.create_user("ana", password="x" * 20))  # nosec B106
    monkeypatch.setattr(
        catalogos,
        "sondar_lectura",
        lambda **k: catalogos.Disponible(controlador=f"{catalogos.PREFIJO_MDBTOOLS}/usr/bin"),
    )
    cuerpo = client.get(reverse("documents:catalogo_a_excel")).content.decode()
    assert 'name="ruta"' in cuerpo or 'type="file"' in cuerpo
