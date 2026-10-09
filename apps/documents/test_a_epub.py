"""«Hacer un libro EPUB»: de PDF, Word, HTML o Markdown a un EPUB 3 que abre en un lector.

## El oráculo

- **EPUBCheck** (W3C, BSD-3), el validador de referencia: `@pytest.mark.oraculo`, con la ruta
  del `.jar` en `AEROCONVERT_EPUBCHECK` y Java en el `PATH`. Procedimiento en
  `docs/PRUEBAS_CON_ORACULO.md`.
- Sin él, en el CI, se lee el libro con **otros lectores** que no lo escribieron: `zipfile` mira el
  contenedor y `defusedxml` el XML. Y el contenido se compara con un **texto de resultado
  conocido** escrito aquí, no con lo que devuelva la función.
"""

from __future__ import annotations

import hashlib
import os
import shutil
import subprocess  # nosec B404 - EPUBCheck, con lista de argumentos
import zipfile
from datetime import UTC, datetime
from pathlib import Path

import pytest
from defusedxml import ElementTree
from django.contrib.auth import get_user_model
from django.urls import reverse

from apps.documents import a_epub, tarea
from apps.documents.composicion import ComposicionInvalida

XHTML = "{http://www.w3.org/1999/xhtml}"
OPF = "{http://www.idpf.org/2007/opf}"

INFORME = """Texto antes del primer título.

# Introducción

Un párrafo con **negrita**, *cursiva* y `código`.

- uno
- dos

## Alcance

| Punto | Este |
| --- | --- |
| PC1 | 414785.848 |

# Resultados

> Una cita.

<script>alert(1)</script>
"""


def _libro(ruta: Path) -> dict:
    """Lo que un lector ve del EPUB, leído sin el módulo que lo escribió."""
    with zipfile.ZipFile(ruta) as z:
        primera = z.infolist()[0]
        opf = ElementTree.fromstring(z.read("OEBPS/content.opf"))
        columna = [r.get("idref") for r in opf.iter(f"{OPF}itemref")]
        por_id = {i.get("id"): i.get("href") for i in opf.iter(f"{OPF}item")}
        capitulos = [ElementTree.fromstring(z.read("OEBPS/" + por_id[i])) for i in columna]
        nav = ElementTree.fromstring(z.read("OEBPS/nav.xhtml"))
        return {
            "primera": primera,
            "mimetype": z.read("mimetype"),
            "titulo": opf.find(".//{http://purl.org/dc/elements/1.1/}title").text,
            "h1": [c.find(f".//{XHTML}h1").text for c in capitulos],
            "indice": [a.text for a in nav.iter(f"{XHTML}a")],
            "textos": ["".join(c.itertext()) for c in capitulos],
            "capitulos": capitulos,
        }


class TestElLibro:
    def test_contenedor_como_lo_exige_un_lector(self, tmp_path):
        destino = tmp_path / "informe.epub"
        a_epub.markdown_a_epub(INFORME, destino, titulo="Informe")
        libro = _libro(destino)
        assert libro["primera"].filename == "mimetype"
        assert libro["primera"].compress_type == zipfile.ZIP_STORED
        assert libro["mimetype"] == b"application/epub+zip"
        assert libro["titulo"] == "Informe"

    def test_un_capitulo_por_titulo_de_primer_nivel_y_el_indice_los_nombra(self, tmp_path):
        destino = tmp_path / "informe.epub"
        a_epub.markdown_a_epub(INFORME, destino, titulo="Informe")
        libro = _libro(destino)
        # Lo de antes del primer título no se pierde: va a un capítulo con el título del libro.
        assert libro["indice"] == ["Informe", "Introducción", "Resultados"]
        assert libro["h1"][1:] == ["Introducción", "Resultados"]
        assert "Texto antes del primer título." in libro["textos"][0]

    def test_formato_listas_y_tablas_llegan(self, tmp_path):
        destino = tmp_path / "informe.epub"
        a_epub.markdown_a_epub(INFORME, destino, titulo="Informe")
        intro = _libro(destino)["capitulos"][1]
        assert intro.find(f".//{XHTML}strong").text == "negrita"
        assert intro.find(f".//{XHTML}em").text == "cursiva"
        assert [li.text for li in intro.iter(f"{XHTML}li")] == ["uno", "dos"]
        assert intro.find(f".//{XHTML}h2").text == "Alcance"
        celdas = [td.text for td in intro.iter(f"{XHTML}td")]
        assert celdas == ["PC1", "414785.848"]

    def test_el_html_incrustado_llega_como_texto_y_no_como_etiqueta(self, tmp_path):
        destino = tmp_path / "informe.epub"
        a_epub.markdown_a_epub(INFORME, destino, titulo="Informe")
        resultados = _libro(destino)["capitulos"][2]
        assert resultados.find(f".//{XHTML}script") is None
        assert "<script>alert(1)</script>" in "".join(resultados.itertext())

    def test_sin_titulos_es_un_solo_capitulo(self, tmp_path):
        destino = tmp_path / "nota.epub"
        a_epub.markdown_a_epub("Solo un párrafo.", destino, titulo="Nota")
        assert _libro(destino)["indice"] == ["Nota"]

    def test_vacio_no_deja_archivo(self, tmp_path):
        destino = tmp_path / "vacio.epub"
        with pytest.raises(ComposicionInvalida, match="No hay texto"):
            a_epub.markdown_a_epub("   \n", destino, titulo="x")
        assert not destino.exists()
        assert not list(tmp_path.glob("*.parcial"))

    def test_mismo_texto_mismo_libro(self, tmp_path):
        """Con fecha e identificador fijos, los bytes no dependen de nada más."""
        momento = datetime(2026, 10, 9, tzinfo=UTC)
        a = tmp_path / "a.epub"
        b = tmp_path / "b.epub"
        for destino in (a, b):
            a_epub.markdown_a_epub(
                INFORME, destino, titulo="I", ahora=momento, identificador="urn:uuid:fijo"
            )
        assert [(i.filename, i.CRC) for i in zipfile.ZipFile(a).infolist()] == [
            (i.filename, i.CRC) for i in zipfile.ZipFile(b).infolist()
        ]


class TestLaTarea:
    def test_desde_word(self, tmp_path):
        from docx import Document

        documento = Document()
        documento.add_heading("Capítulo uno", level=1)
        documento.add_paragraph("Texto del primero.")
        documento.add_heading("Capítulo dos", level=1)
        documento.add_paragraph("Texto del segundo.")
        origen = tmp_path / "manual.docx"
        documento.save(origen)
        antes = hashlib.sha256(origen.read_bytes()).hexdigest(), origen.stat().st_mtime_ns

        parcial = tmp_path / "salida.parcial.epub"
        informe = tarea.ejecutar(
            "a_epub", {"entradas": [{"ruta": str(origen)}], "opciones": {}}, parcial
        )
        assert informe.get("codigo") is None, informe
        libro = _libro(parcial)
        assert libro["titulo"] == "manual"
        assert libro["indice"] == ["Capítulo uno", "Capítulo dos"]
        assert "Texto del segundo." in libro["textos"][1]
        assert (hashlib.sha256(origen.read_bytes()).hexdigest(), origen.stat().st_mtime_ns) == antes

    def test_un_pdf_sin_texto_lo_dice_y_no_deja_libro(self, tmp_path):
        from pypdf import PdfWriter

        escritor = PdfWriter()
        escritor.add_blank_page(width=595, height=842)
        origen = tmp_path / "escaneo.pdf"
        with origen.open("wb") as f:
            escritor.write(f)
        parcial = tmp_path / "salida.parcial.epub"
        informe = tarea.ejecutar(
            "a_epub", {"entradas": [{"ruta": str(origen)}], "opciones": {}}, parcial
        )
        assert informe["desenlace"] == "sin-texto-que-sacar"
        assert not parcial.exists()


@pytest.mark.django_db
class TestLaPantalla:
    @pytest.fixture
    def sesion(self, client, tmp_path, settings):
        settings.RAICES_PERMITIDAS = str(tmp_path)
        settings.CARPETA_DE_TRABAJO = str(tmp_path / "trabajo")
        client.force_login(get_user_model().objects.create_user("lee"))
        return client

    def test_pide_sesion(self, client):
        assert client.get(reverse("documents:a_epub")).status_code == 302

    def test_rechaza_lo_que_no_sabe_leer(self, sesion, tmp_path):
        origen = tmp_path / "libro.epub"
        origen.write_bytes(b"PK")
        cuerpo = sesion.post(reverse("documents:a_epub"), {"ruta": str(origen)}).content.decode()
        assert "no se hace un libro" in cuerpo

    def test_de_extremo_a_extremo_desde_markdown(self, sesion, tmp_path):
        from apps.jobs import despachador
        from apps.jobs.models import ConversionJob

        origen = tmp_path / "informe.md"
        origen.write_text(INFORME, encoding="utf-8")
        respuesta = sesion.post(
            reverse("documents:a_epub"),
            {"ruta": str(origen), "titulo": "Informe de vuelo", "autor": "Topografía"},
        )
        assert respuesta.status_code == 302 and "/trabajos/" in respuesta["Location"]
        assert despachador.procesar_una_vez() == 1
        trabajo = ConversionJob.objects.latest("created_at")
        assert trabajo.status == "done", trabajo.reason_detail
        salida = Path(trabajo.output_path)
        assert salida.suffix == ".epub"
        assert _libro(salida)["titulo"] == "Informe de vuelo"


@pytest.mark.oraculo
def test_epubcheck_lo_da_por_valido(tmp_path):
    jar = os.environ.get("AEROCONVERT_EPUBCHECK", "")
    java = shutil.which("java")
    if not (jar and Path(jar).is_file() and java):
        pytest.skip("Hace falta EPUBCheck (AEROCONVERT_EPUBCHECK) y Java.")
    destino = tmp_path / "informe.epub"
    a_epub.markdown_a_epub(INFORME, destino, titulo="Informe", autor="Topografía")
    resultado = subprocess.run(  # nosec B603 - rutas propias, sin shell
        [java, "-jar", jar, str(destino)], capture_output=True, text=True, timeout=120
    )
    # EPUBCheck sale con 0 sin errores, pero la regla 1 dice que el código no basta: se lee
    # además su frase final.
    assert resultado.returncode == 0, resultado.stdout + resultado.stderr
    assert "No errors or warnings detected" in resultado.stdout, resultado.stdout
